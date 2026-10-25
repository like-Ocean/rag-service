from io import BytesIO
from pathlib import Path
from service.pdf_parser import parse_pdf

MIME_TO_EXTENSION = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
}


class UnsupportedDocumentFormatError(ValueError):
    pass


def _docx_paragraph_segments(paragraph) -> list[str]:
    from docx.oxml.ns import qn

    segments = [""]
    for run in paragraph.runs:
        for element in run._r.iter():
            if element.tag == qn("w:t"):
                segments[-1] += element.text or ""
            elif element.tag == qn("w:tab"):
                segments[-1] += "\t"
            elif element.tag == qn("w:br"):
                if element.get(qn("w:type")) == "page":
                    segments.append("")
                else:
                    segments[-1] += "\n"
            elif element.tag == qn("w:lastRenderedPageBreak"):
                segments.append("")
    if not paragraph.runs and paragraph.text:
        segments[-1] = paragraph.text
    return segments


def _parse_docx(document_bytes: bytes) -> list[tuple[int, str]]:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(document_bytes))
    page_blocks: dict[int, list[str]] = {1: []}
    page_no = 1

    def advance_page() -> None:
        nonlocal page_no
        page_no += 1
        page_blocks.setdefault(page_no, [])

    for block in document.element.body.iterchildren():
        if isinstance(block, CT_P):
            paragraph = Paragraph(block, document)
            if paragraph.paragraph_format.page_break_before and page_blocks[page_no]:
                advance_page()
            for segment_index, segment in enumerate(_docx_paragraph_segments(paragraph)):
                if segment_index:
                    advance_page()
                if segment.strip():
                    page_blocks[page_no].append(segment.strip())
        elif isinstance(block, CT_Tbl):
            table = Table(block, document)
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    page_blocks[page_no].append(row_text)

    return [
        (number, "\n\n".join(blocks))
        for number, blocks in sorted(page_blocks.items())
        if blocks
    ]


def _parse_pptx(document_bytes: bytes) -> list[tuple[int, str]]:
    from pptx import Presentation

    presentation = Presentation(BytesIO(document_bytes))
    slides: list[tuple[int, str]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        blocks.append(row_text)
            elif shape.has_text_frame and shape.text.strip():
                blocks.append(shape.text.strip())
        slides.append((slide_number, "\n\n".join(blocks)))
    return slides


def parse_document(
    document_bytes: bytes, filename: str | None = None,
    mime_type: str | None = None
) -> list[tuple[int, str]]:
    extension = Path(filename or "").suffix.casefold()
    if not extension and mime_type:
        extension = MIME_TO_EXTENSION.get(mime_type.split(";", 1)[0].strip().casefold(), "")

    if extension == ".pdf":
        return parse_pdf(document_bytes)
    if extension == ".docx":
        return _parse_docx(document_bytes)
    if extension == ".pptx":
        return _parse_pptx(document_bytes)
    if extension == ".doc":
        raise UnsupportedDocumentFormatError(
            "Legacy .doc files are not supported yet; convert to .docx or PDF."
        )
    raise UnsupportedDocumentFormatError(
        "Supported formats: PDF, DOCX and PPTX. Provide the original filename or MIME type."
    )
