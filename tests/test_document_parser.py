from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from service.document_parser import UnsupportedDocumentFormatError, parse_document


def test_docx_extracts_paragraphs_tables_and_explicit_page_breaks():
    document = Document()
    document.add_paragraph("Первый раздел")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Термин"
    table.cell(0, 1).text = "Определение"
    document.add_page_break()
    document.add_paragraph("Второй раздел")
    stream = BytesIO()
    document.save(stream)

    pages = parse_document(stream.getvalue(), "study-guide.docx")

    assert pages[0][0] == 1
    assert "Первый раздел" in pages[0][1]
    assert "Термин | Определение" in pages[0][1]
    assert pages[-1] == (2, "Второй раздел")


def test_pptx_extracts_text_and_tables_by_slide_number():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "Слайд первый"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    stream = BytesIO()
    presentation.save(stream)

    pages = parse_document(
        stream.getvalue(),
        "presentation.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    assert pages[0][0] == 1
    assert "Слайд первый" in pages[0][1]
    assert "Показатель | Значение" in pages[0][1]
    assert pages[1] == (2, "")


def test_legacy_doc_is_reported_as_unsupported():
    with pytest.raises(UnsupportedDocumentFormatError, match="Legacy .doc"):
        parse_document(b"legacy word", "old-file.doc")


def test_document_type_can_be_inferred_from_mime_when_filename_is_missing():
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    stream = BytesIO()
    presentation.save(stream)

    assert parse_document(
        stream.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ) == [(1, "")]
